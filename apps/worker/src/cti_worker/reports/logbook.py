"""Logbook Threat Information (Excel, tiap 2 minggu) -- gantiin
`ScraperNews/supportFile/logbook.py`.

Daftar berita GLOBAL/APAC 14 hari terakhir dimasukkan ke templat dokumen resmi
(`templates/logbook_template.xlsx`, disalin dari `supportFile/templateLogbook.xlsx`),
lengkap dengan blok tanda tangan.

Perbedaan dari skrip lama:
  - Tanggal berbahasa Indonesia dari tabel bulan statis, BUKAN `GoogleTranslator`
    (`deep_translator` = satu panggilan jaringan ke Google PER tanggal; ratusan
    artikel = ratusan panggilan, dan laporan gagal kalau Google menolak).
  - Data dari tabel `articles` (skrip lama: koleksi Mongo `articles` juga -- sama),
    kategori dari `Article.news_type`. Cabang "APT" di skrip lama tidak pernah
    tercapai (`... not in news_type` sesudah `in news_type`), jadi kategori APT gak
    ikut -- dipertahankan.
  - Nama + jabatan penandatangan dari setting (`WorkerSettings.logbook_*`), bukan
    hardcoded.
  - Skrip lama MEMBUAT file lalu berhenti (baris `send_file` dikomentari) --
    laporannya cuma ada di disk server. Sekarang dikirim ke Telegram (topik `logbook`).
  - Mengembalikan bytes (tanpa file sementara di CWD, yang bercampur antar run).

Yang dipertahankan APA ADANYA: urutan operasi openpyxl (`insert_rows` lalu isi sel),
border, selang-seling warna, hyperlink "Link to Article".
"""

from __future__ import annotations

import datetime
import io
from collections.abc import Sequence
from dataclasses import dataclass
from importlib import resources

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_MONTHS_ID = (
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
)  # fmt: skip
_MONTHS_EN = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
_CATEGORIES = ("GLOBAL", "APAC")

_thin = Side(style="thin")
_medium = Side(style="medium")
_B_LEFT = Border(top=_thin, right=_thin, left=_medium, bottom=_thin)
_B_LEFT_ONLY = Border(left=_medium)
_B_RIGHT_ONLY = Border(right=_medium)
_B_RIGHT = Border(right=_medium)
_B_THIN = Border(top=_thin, right=_thin, left=_thin, bottom=_thin)
_B_THIN_LR = Border(right=_thin, left=_thin)
_FILL = PatternFill(start_color="D6E1F4", end_color="D6E1F4", fill_type="solid")


@dataclass(frozen=True)
class LogbookEntry:
    category: str
    title: str
    url: str
    posted_on: datetime.date


@dataclass(frozen=True)
class Signatories:
    preparer_name: str
    preparer_title: str
    approver_name: str
    approver_title: str


def format_date_id(day: datetime.date, *, zero_pad: bool = True) -> str:
    """ "26 September 2026". `zero_pad`: header dokumen memakai "05 Maret", baris
    artikel "5 Maret" (sama dengan skrip lama)."""
    d = f"{day.day:02d}" if zero_pad else str(day.day)
    return f"{d} {_MONTHS_ID[day.month - 1]} {day.year}"


def filename_for(start: datetime.date, end: datetime.date) -> str:
    def fmt(d: datetime.date) -> str:
        return f"{d.day:02d}-{_MONTHS_EN[d.month - 1]}-{d.year}"

    return f"ThreatInformation_Logbook_{fmt(start)}_to_{fmt(end)}.xlsx"


def _put(ws, coord: str, border: Border, value: object, fill: bool) -> None:
    cell = ws[coord]
    cell.border = border
    horizontal = "left" if coord.startswith(("E", "G")) else "center"
    cell.alignment = Alignment(horizontal=horizontal, vertical="center")
    try:
        if str(value).startswith("http"):
            cell.hyperlink = str(value)
            cell.value = "Link to Article"
        else:
            cell.value = value
    except AttributeError:
        pass  # sel F tergabung dgn E (`merge_cells`) -> read-only; sama dgn skrip lama
    if fill:
        cell.fill = _FILL


def _sign(ws, coord: str, border: Border, value: str, underline: bool) -> None:
    cell = ws[coord]
    cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.value = value
    cell.font = Font(name="Calibri", size=12, bold=True, underline="single" if underline else None)


def build(
    entries: Sequence[LogbookEntry],
    *,
    start: datetime.date,
    end: datetime.date,
    signatories: Signatories,
) -> bytes:
    template = resources.files("cti_worker.reports").joinpath("templates/logbook_template.xlsx")
    with resources.as_file(template) as path:
        workbook = openpyxl.load_workbook(path)
    ws = workbook[workbook.sheetnames[0]]

    period = f"{format_date_id(start)} - {format_date_id(end)}"
    ws.cell(row=3, column=7).value = f"Tanggal: {format_date_id(end)}"
    ws.cell(row=3, column=7).border = _B_RIGHT_ONLY
    ws.cell(row=4, column=5).value = f"Logbook Threat Information Periode\n{period}"
    ws.cell(row=11, column=5).value = (
        "Laporan ini diberikan setiap dua minggu sekali di mana dalam laporan ini berisi "
        "informasi ancaman siber yang terjadi dalam periode "
        f"{period.replace(' - ', ' \u2013 ')}. Informasi dikategorikan berdasarkan "
        "cakupan geografis dan jenis ancamannya."
    )
    for row in range(1, 5):
        ws.cell(row=row, column=10).border = _B_LEFT_ONLY

    row_no = 13
    while ws.cell(row=row_no, column=2).value is not None:
        row_no += 1

    ws.insert_rows(idx=row_no, amount=len(entries))
    fill = False
    for counter, entry in enumerate(entries, start=1):
        ws.merge_cells(f"E{row_no}:F{row_no}", start_row=row_no)
        _put(ws, f"B{row_no}", _B_LEFT, counter, fill)
        _put(ws, f"C{row_no}", _B_THIN, format_date_id(entry.posted_on, zero_pad=False), fill)
        _put(ws, f"D{row_no}", _B_THIN, entry.category, fill)
        _put(ws, f"E{row_no}", _B_THIN, entry.title, fill)
        _put(ws, f"F{row_no}", _B_THIN, "", False)
        _put(ws, f"G{row_no}", _B_THIN, entry.url, fill)
        _put(ws, f"H{row_no}", _B_THIN, "-", fill)
        _put(ws, f"I{row_no}", _B_RIGHT, "", False)
        fill = not fill
        row_no += 1

    ws.insert_rows(idx=row_no, amount=15)
    for counter, i in enumerate(range(row_no, row_no + 14), start=1):
        _sign(ws, f"B{i}", _B_LEFT_ONLY, "", False)
        _sign(ws, f"I{i}", _B_RIGHT, "", False)
        if counter == 6:
            _sign(ws, f"E{i}", _B_THIN, "Disusun Oleh:", False)
            _sign(ws, f"F{i}", _B_THIN, "Disetujui Oleh:", False)
        if 7 <= counter < 13:
            _sign(ws, f"E{i}", _B_THIN_LR, "", False)
            _sign(ws, f"F{i}", _B_THIN_LR, "", False)
        if counter == 13:
            _sign(ws, f"E{i}", _B_THIN, signatories.preparer_name, True)
            _sign(ws, f"F{i}", _B_THIN, signatories.approver_name, True)
        if counter == 14:
            _sign(ws, f"E{i}", _B_THIN, signatories.preparer_title, False)
            _sign(ws, f"F{i}", _B_THIN, signatories.approver_title, False)
            _sign(ws, f"B{i + 1}", _B_LEFT_ONLY, "", False)
            _sign(ws, f"I{i + 1}", _B_RIGHT, "", False)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def caption(start: datetime.date, end: datetime.date) -> str:
    def fmt(d: datetime.date) -> str:
        return f"{d.day:02d}-{_MONTHS_EN[d.month - 1]}-{d.year}"

    return (
        "\n    === THREAT INFORMATION LOGBOOK REPORT ===\n"
        f"<i>This alert is designed for reporting news scraped throughout {fmt(start)} to "
        f"{fmt(end)}</i>"
    )
