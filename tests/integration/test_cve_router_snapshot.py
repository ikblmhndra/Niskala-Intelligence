"""Snapshot test `routers/cve.py` -- Fase 7.6, PILOT router buat pola
"snapshot test tiap endpoint" (20 endpoint, terbanyak dari 27 router).

Beda dari test integrasi lain di suite ini (yang manggil fungsi
service/repo LANGSUNG lewat `async_db_session`): di sini request beneran
lewat HTTP (`api_client`, ASGI transport ke `create_app()`) -- routing,
auth dependency (`require_auth`/`effective_client_id`), query/body
validation Pydantic, DAN `response_model` serialization semua kepake,
gak cuma logic service. `api_client`/`api_session` (lihat
`tests/integration/conftest.py`) ngejaga isolasi per-test walau
endpoint-nya beneran manggil `session.commit()`.

Snapshot (`syrupy`) nangkep BENTUK response biar drift (field ilang/
ganti nama/ganti tipe) ketahuan otomatis pas refactor nanti -- bukan
buat re-verifikasi correctness (itu udah dites integration test service-
layer yang ada). Field yang GAK MUNGKIN deterministik lintas run
(`id` auto-increment Postgres gak ke-reset rollback transaksi test,
`last_updated`/`acknowledge_time` timestamp wall-clock, `ticket_id`
nempel bulan-tahun kalender asli) dinormalisasi lewat matcher
`path_type` -- diganti placeholder tipe, bukan dihapus dari
perbandingan, jadi tetep ketahuan kalau field itu ILANG atau GANTI TIPE.

Endpoint yang manggil API eksternal (`cisa-lookup`/`epss-lookup`/
`exploit-lookup` -- Grup B; `draft-email` -- Grup C, LLM+Graph) di-mock
di titik yang SAMA persis kayak test service-layer yang udah ada
(`test_cve_lookup_service.py`/`test_cve_email_service.py`) -- gak ada
panggilan jaringan asli dari test ini."""

from __future__ import annotations

import datetime
import io
from collections.abc import Callable
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl
import pytest
from cti_api.services import cve_email as cve_email_service
from cti_api.services import cve_lookup as cve_lookup_service
from cti_core.db.models.cve import CveTracker
from cti_core.db.models.techstack import TechStackEntry
from cti_core.db.repositories.auth import AsyncClientRepo
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession
from syrupy.assertion import SnapshotAssertion
from syrupy.matchers import path_type

pytestmark = pytest.mark.asyncio

_NORMALIZE = path_type(
    {
        r"(.*\.)?id$": (int,),
        r"(.*\.)?last_updated$": (str,),
        r"(.*\.)?acknowledge_time$": (str,),
        r"(.*\.)?ticket_id$": (str,),
    },
    regex=True,
    strict=False,
)


async def _seed_two_cves(session: AsyncSession, client_id: str = "default") -> None:
    await AsyncClientRepo(session).ensure_default()
    session.add_all(
        [
            CveTracker(
                cve_id="CVE-2026-7001",
                client_id=client_id,
                tech="WordPress",
                summary="Stored XSS in a WordPress plugin",
                cve_score=9.8,
                cve_severity="CRITICAL",
                published=datetime.date(2026, 9, 1),
                detected_on=datetime.datetime(2026, 9, 1, tzinfo=datetime.UTC),
                cisa_kev=True,
            ),
            CveTracker(
                cve_id="CVE-2026-7002",
                client_id=client_id,
                tech="nginx",
                summary="Config parsing bug",
                cve_score=3.1,
                cve_severity="LOW",
                published=datetime.date(2026, 9, 5),
                detected_on=datetime.datetime(2026, 9, 5, tzinfo=datetime.UTC),
            ),
        ]
    )
    await session.flush()


async def test_list_cves(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.get("/api/cve", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot(matcher=_NORMALIZE)


async def test_cve_stats(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.get("/api/cve/stats", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_tech_list(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.get("/api/cve/tech-list", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_set_false_positive(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.post("/api/cve/CVE-2026-7001/false-positive", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_remove_false_positive(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    await api_client.post("/api/cve/CVE-2026-7001/false-positive", headers=auth_header())
    resp = await api_client.delete("/api/cve/CVE-2026-7001/false-positive", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_bulk_false_positive(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.post(
        "/api/cve/bulk-false-positive",
        headers=auth_header(),
        json={"cve_ids": ["CVE-2026-7001", "CVE-2026-7002"]},
    )
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_purge_orphaned_dry_run(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.post(
        "/api/cve/purge-orphaned", headers=auth_header(), json={"dry_run": True}
    )
    assert resp.status_code == 200
    body = resp.json()
    body["cve_ids"] = sorted(body["cve_ids"])
    assert body == snapshot


async def test_cisa_kev_lookup(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    fake_catalog = {"CVE-2026-7001": {"vendorProject": "wp", "dateAdded": "2026-01-01"}}
    with patch.object(cve_lookup_service, "_fetch_cisa_kev", AsyncMock(return_value=fake_catalog)):
        resp = await api_client.post("/api/cve/cisa-lookup", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_epss_lookup(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    fake_scores = {"CVE-2026-7001": {"epss": 0.5, "percentile": 0.8, "date": "2026-09-19"}}
    with patch.object(
        cve_lookup_service, "_fetch_epss_scores", AsyncMock(return_value=fake_scores)
    ):
        resp = await api_client.post("/api/cve/epss-lookup", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_bulk_exploit_lookup(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    with (
        patch.object(cve_lookup_service, "_lookup_exploitdb", AsyncMock(return_value=[])),
        patch("asyncio.sleep", AsyncMock()),
    ):
        resp = await api_client.post("/api/cve/exploit-lookup", headers=auth_header())
    assert resp.status_code == 200
    body = resp.json()
    # `run_exploit_db_bulk_lookup()` urut dari `list_true_positive_cve_ids()`
    # -> `list_all_distinct_cve_ids()`, yang query-nya `SELECT DISTINCT`
    # TANPA `ORDER BY` -- Postgres gak janji urutan row buat DISTINCT
    # tanpa itu, dan urutannya bisa geser tergantung query plan (yang
    # kepengaruh histori churn tabel dari test LAIN pas full-suite run,
    # gak kejadian kalau file ini dijalanin sendirian). Disortir di sini
    # biar snapshot stabil -- normalisasi test, bukan ubah query service
    # (pre-existing dari Fase 7.3 Bagian 4, ketauan gak sengaja pas
    # verifikasi full-suite Fase 7.6, sama pola kayak `cluster_service`'s
    # `sources` set-ordering).
    body["details"] = sorted(body["details"], key=lambda d: d["cve_id"])
    assert body == snapshot


async def test_single_exploit_lookup(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    fake_exploits = [
        {
            "edb_id": "9",
            "title": "t",
            "posted_on": "2026-01-01",
            "url": "https://www.exploit-db.com/exploits/9",
            "source": "exploit-db",
            "type": "exploit",
        }
    ]
    with patch.object(
        cve_lookup_service, "_lookup_exploitdb", AsyncMock(return_value=fake_exploits)
    ):
        resp = await api_client.post("/api/cve/CVE-2026-7001/exploit-lookup", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_export_cves(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    """Response binary (.xlsx) -- di-parse lewat openpyxl, snapshot
    METADATA-nya (header kolom + jumlah baris), bukan byte mentah
    (nama file/timestamp internal xlsx non-deterministik)."""
    await _seed_two_cves(api_session)
    resp = await api_client.get("/api/cve/export", headers=auth_header())
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    data_row_count = ws.max_row - 1
    assert {"headers": header_row, "data_row_count": data_row_count} == snapshot


async def test_next_ticket_id(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await AsyncClientRepo(api_session).ensure_default()
    resp = await api_client.get("/api/cve/next-ticket-id", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot(matcher=_NORMALIZE)


async def test_ack_statuses_empty(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.get("/api/cve/ack-statuses", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_draft_email(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    """LLM (`_llm_cve_validator`/`_dedup_mitigation`/`_dedup_risk_context`)
    + Graph draft (`create_graph_draft`) di-mock, sama titik persis kayak
    `test_cve_email_service.py` -- endpoint ini SENGAJA belum live-tested
    lawan API asli (lihat docstring `cve_email.py`, keputusan user Grup C,
    cost token + bikin draft asli)."""
    await _seed_two_cves(api_session)
    fake_validation = {
        "product_name": "WordPress",
        "summary": "Attackers can execute arbitrary PHP code.",
        "risk_context": "Full compromise possible if unpatched.",
        "primary_remediation_action": [{"version_branch": "< 6.5", "fixed_version": "6.5"}],
        "alternative_remediation_action": ["Disable the plugin"],
    }
    with (
        patch.object(
            cve_email_service, "_llm_cve_validator", MagicMock(return_value=fake_validation)
        ),
        patch.object(
            cve_email_service,
            "_dedup_mitigation",
            MagicMock(return_value={"alternative_remediation_action": ["Deduped control"]}),
        ),
        patch.object(
            cve_email_service,
            "_dedup_risk_context",
            MagicMock(return_value="<strong>Risky</strong>"),
        ),
        patch.object(
            cve_email_service, "create_graph_draft", MagicMock(return_value="graph-msg-123")
        ),
    ):
        resp = await api_client.post(
            "/api/cve/draft-email",
            headers=auth_header(),
            json={"cve_ids": ["CVE-2026-7001"], "ticket_id": "CTI-2026-09-999"},
        )
    assert resp.status_code == 200
    assert resp.json() == snapshot(matcher=_NORMALIZE)


async def test_get_cve_ticket_empty(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.get("/api/cve/CVE-2026-7001/ticket", headers=auth_header())
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_update_cve_ticket_then_get(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    put_resp = await api_client.put(
        "/api/cve/CVE-2026-7001/ticket",
        headers=auth_header(),
        json={
            "affected_asset": "web-01",
            "asset_owner": "IT Team",
            "remediation_status": "in_progress",
            "remediation_date_plan": "2026-09-15",
        },
    )
    assert put_resp.status_code == 200
    assert put_resp.json() == snapshot(name="put_response")

    get_resp = await api_client.get("/api/cve/CVE-2026-7001/ticket", headers=auth_header())
    assert get_resp.status_code == 200
    assert get_resp.json() == snapshot(name="get_after_put", matcher=_NORMALIZE)


async def test_bulk_acknowledge(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.post(
        "/api/cve/bulk-acknowledge",
        headers=auth_header(),
        json={"cve_ids": ["CVE-2026-7001", "CVE-2026-7002"], "analyst_name": "Analyst A"},
    )
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_acknowledge(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    resp = await api_client.post(
        "/api/cve/CVE-2026-7001/acknowledge",
        headers=auth_header(),
        json={"analyst_name": "Analyst A"},
    )
    assert resp.status_code == 200
    assert resp.json() == snapshot


async def test_prioritize_cves(
    api_client: AsyncClient,
    api_session: AsyncSession,
    auth_header: Callable[..., dict[str, str]],
    snapshot: SnapshotAssertion,
) -> None:
    await _seed_two_cves(api_session)
    api_session.add(TechStackEntry(name="WordPress", client_id="default"))
    await api_session.flush()
    resp = await api_client.get(
        "/api/cve/prioritize", headers=auth_header(), params={"cves": "CVE-2026-7001,CVE-2026-7002"}
    )
    assert resp.status_code == 200
    assert resp.json() == snapshot
