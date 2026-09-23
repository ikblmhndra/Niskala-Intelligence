"""Integration test `cti_api.services.cve_export.export_cves_to_excel` --
Postgres REAL (testcontainers) + openpyxl beneran (bukan mock), template
`.xlsx` asli. Fase 7.4 Grup C (2026-09-23)."""

from __future__ import annotations

import datetime
import io

import openpyxl
import pytest
from cti_api.services import cve_export as svc
from cti_core.db.models.cve import CveTracker
from cti_core.db.repositories.auth import AsyncClientRepo
from cti_core.db.repositories.cve import AsyncCveFalsePositiveRepo, AsyncCveTicketRepo
from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.asyncio


async def _ensure_clients(session: AsyncSession) -> None:
    repo = AsyncClientRepo(session)
    await repo.ensure_default()
    if await repo.get("acme") is None:
        await repo.create(client_id="acme", name="Acme", countries=[])


def _load(data: bytes) -> tuple[object, object]:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["CVE Ticket Tracker"]
    return wb, ws


async def test_export_writes_cve_and_ticket_fields(async_db_session: AsyncSession) -> None:
    await _ensure_clients(async_db_session)
    async_db_session.add(
        CveTracker(
            cve_id="CVE-2026-0001",
            client_id="default",
            tech="WordPress",
            cve_severity="CRITICAL",
            cve_score=9.8,
            published=datetime.date(2026, 9, 1),
            solutions="Upgrade to 6.5",
            active_exploitation=True,
        )
    )
    await async_db_session.flush()
    ticket = await AsyncCveTicketRepo(async_db_session).upsert(
        cve_id="CVE-2026-0001",
        client_id="default",
        affected_asset="web-01",
        owner_email="owner@example.com",
        escalation_required=True,
        remediation_date_plan=datetime.date(2026, 10, 1),
    )

    data = await svc.export_cves_to_excel(async_db_session, client_id="default")
    _, ws = _load(data)

    assert ws.cell(row=16, column=2).value == ticket.ticket_id
    assert ws.cell(row=16, column=3).value == "CVE-2026-0001"
    assert ws.cell(row=16, column=4).value == "2026-09-01"  # cve_reported_date derived
    assert ws.cell(row=16, column=5).value == "CRITICAL"
    assert ws.cell(row=16, column=6).value == 9.8
    assert ws.cell(row=16, column=7).value == "web-01"
    assert ws.cell(row=16, column=11).value == "owner@example.com"
    assert ws.cell(row=16, column=15).value == "2026-10-01"
    assert ws.cell(row=16, column=18).value == "Yes"  # escalation_required bool -> Yes/No


async def test_export_falls_back_to_cve_fields_when_no_ticket(
    async_db_session: AsyncSession,
) -> None:
    await _ensure_clients(async_db_session)
    async_db_session.add(
        CveTracker(
            cve_id="CVE-2026-0002",
            client_id="default",
            tech="nginx",
            solutions="Patch to 1.27",
            active_exploitation=True,
        )
    )
    await async_db_session.flush()

    data = await svc.export_cves_to_excel(async_db_session, client_id="default")
    _, ws = _load(data)

    assert ws.cell(row=16, column=2).value in ("", None)  # ticket_id kosong, gak ada ticket
    assert ws.cell(row=16, column=7).value == "nginx"  # affected_asset fallback ke cve.tech
    assert ws.cell(row=16, column=9).value == "Patch to 1.27"  # fixed_version fallback solutions
    assert ws.cell(row=16, column=13).value == "Yes"  # active_exploitation fallback dari cve


async def test_export_excludes_false_positives_by_default(async_db_session: AsyncSession) -> None:
    await _ensure_clients(async_db_session)
    async_db_session.add_all(
        [
            CveTracker(cve_id="CVE-2026-0003", client_id="default", tech="A"),
            CveTracker(cve_id="CVE-2026-0004", client_id="default", tech="B"),
        ]
    )
    await async_db_session.flush()
    await AsyncCveFalsePositiveRepo(async_db_session).mark("CVE-2026-0004", "default")

    data = await svc.export_cves_to_excel(async_db_session, client_id="default")
    _, ws = _load(data)

    cve_ids = [
        ws.cell(row=r, column=3).value for r in range(16, 18) if ws.cell(row=r, column=3).value
    ]
    assert cve_ids == ["CVE-2026-0003"]


async def test_export_scoped_to_client_id(async_db_session: AsyncSession) -> None:
    """Regression buat bug yang diperbaiki -- legacy `export_cves_to_
    excel()` gak pernah nge-scope client_id sama sekali. Export client
    "acme" HARUS gak narik CVE client "default"."""
    await _ensure_clients(async_db_session)
    async_db_session.add_all(
        [
            CveTracker(cve_id="CVE-2026-0005", client_id="default", tech="A"),
            CveTracker(cve_id="CVE-2026-0006", client_id="acme", tech="B"),
        ]
    )
    await async_db_session.flush()

    data = await svc.export_cves_to_excel(async_db_session, client_id="acme")
    _, ws = _load(data)

    assert ws.cell(row=16, column=3).value == "CVE-2026-0006"
    assert ws.cell(row=17, column=3).value is None


async def test_export_empty_result_produces_no_data_rows(async_db_session: AsyncSession) -> None:
    await _ensure_clients(async_db_session)
    data = await svc.export_cves_to_excel(async_db_session, client_id="default")
    _, ws = _load(data)
    assert ws.cell(row=16, column=3).value is None
