"""`cti_worker.reports.logbook` -- logbook Excel dua-mingguan (Fase 10.E).

Dibaca balik dengan openpyxl: yang dikunci isi sel (header, baris artikel,
hyperlink, blok tanda tangan), bukan piksel.
"""

from __future__ import annotations

import datetime
import io

import openpyxl
from cti_worker.reports import logbook

START, END = datetime.date(2026, 9, 12), datetime.date(2026, 9, 26)
SIGN = logbook.Signatories("Ani Preparer", "Threat Intel", "Budi Approver", "Security Lead")


def entries(n: int) -> list[logbook.LogbookEntry]:
    return [
        logbook.LogbookEntry(
            "GLOBAL" if i % 2 == 0 else "APAC",
            f"Judul berita {i}",
            f"https://example.com/{i}",
            datetime.date(2026, 9, 5 + i),
        )
        for i in range(n)
    ]


def load(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def test_dates_render_in_indonesian_with_and_without_zero_padding() -> None:
    assert logbook.format_date_id(datetime.date(2026, 3, 5)) == "05 Maret 2026"
    assert logbook.format_date_id(datetime.date(2026, 3, 5), zero_pad=False) == "5 Maret 2026"
    assert logbook.format_date_id(datetime.date(2026, 12, 31)) == "31 Desember 2026"
    assert logbook.format_date_id(datetime.date(2026, 8, 17)) == "17 Agustus 2026"


def test_filename_uses_english_month_abbreviations_regardless_of_locale() -> None:
    assert logbook.filename_for(START, END) == (
        "ThreatInformation_Logbook_12-Sep-2026_to_26-Sep-2026.xlsx"
    )


def test_header_carries_the_period_and_dates() -> None:
    ws = load(logbook.build(entries(1), start=START, end=END, signatories=SIGN))

    assert ws["G3"].value == "Tanggal: 26 September 2026"
    assert ws["E4"].value == (
        "Logbook Threat Information Periode\n12 September 2026 - 26 September 2026"
    )
    assert "12 September 2026 \u2013 26 September 2026" in ws["E11"].value


def test_article_rows_start_below_the_template_header_and_link_to_the_article() -> None:
    ws = load(logbook.build(entries(3), start=START, end=END, signatories=SIGN))

    assert [ws.cell(row=14, column=c).value for c in (2, 3, 4, 5)] == [
        1, "5 September 2026", "GLOBAL", "Judul berita 0",
    ]  # fmt: skip
    assert (
        ws["G14"].value == "Link to Article"
        and ws["G14"].hyperlink.target == "https://example.com/0"
    )
    assert ws["H14"].value == "-"
    assert [ws.cell(row=r, column=2).value for r in (14, 15, 16)] == [1, 2, 3]
    assert ws["D15"].value == "APAC"


def test_rows_alternate_fill_starting_unfilled() -> None:
    ws = load(logbook.build(entries(3), start=START, end=END, signatories=SIGN))

    fills = [ws.cell(row=r, column=3).fill.start_color.rgb for r in (14, 15, 16)]
    assert fills[0] != fills[1] and fills[0] == fills[2]
    assert fills[1].endswith("D6E1F4")


def test_signature_block_follows_the_last_article_with_configured_names() -> None:
    ws = load(logbook.build(entries(2), start=START, end=END, signatories=SIGN))

    cells = {c.value: c.coordinate for row in ws.iter_rows() for c in row if c.value}
    assert cells["Disusun Oleh:"] == "E21" and cells["Disetujui Oleh:"] == "F21"  # 14+2 baris + 5
    assert cells["Ani Preparer"] == "E28" and cells["Budi Approver"] == "F28"
    assert cells["Threat Intel"] == "E29" and cells["Security Lead"] == "F29"
    assert ws["E28"].font.underline == "single" and ws["E28"].font.bold


def test_an_empty_period_still_produces_a_valid_workbook() -> None:
    ws = load(logbook.build([], start=START, end=END, signatories=SIGN))

    assert ws["G3"].value == "Tanggal: 26 September 2026"
    assert ws["B14"].value is None  # tidak ada baris artikel


def test_caption_names_the_period() -> None:
    text = logbook.caption(START, END)

    assert "THREAT INFORMATION LOGBOOK REPORT" in text and "12-Sep-2026 to 26-Sep-2026" in text
